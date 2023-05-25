import pandas as pd
from docx import Document


df = pd.read_csv('output.csv')

document = Document()
table = document.add_table(rows=df.shape[0]+1, cols=df.shape[1])


for i, column_name in enumerate(df.columns):#adding heads
    table.cell(0, i).text = column_name

for i in range(df.shape[0]):
    for j in range(df.shape[1]):
        table.cell(i+1, j).text = str(df.iloc[i, j])


document.save('output.docx')

print("Data exported to Word successfully!")
